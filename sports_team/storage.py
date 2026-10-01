import sqlite3
from typing import List, Optional
from sports_team.models import Player, Team, Match
import docx
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

DB_NAME = "sports_analytics.db"


def get_connection(db_path: str = DB_NAME) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def init_db(db_path: str = DB_NAME):
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS teams (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            wins INTEGER DEFAULT 0,
            draws INTEGER DEFAULT 0,
            losses INTEGER DEFAULT 0,
            goals_for INTEGER DEFAULT 0,
            goals_against INTEGER DEFAULT 0
        );
        """)

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS players (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            position TEXT NOT NULL,
            matches INTEGER DEFAULT 0,
            goals INTEGER DEFAULT 0,
            assists INTEGER DEFAULT 0,
            penalty_minutes INTEGER DEFAULT 0,
            team_id INTEGER,
            FOREIGN KEY (team_id) REFERENCES teams (id) ON DELETE SET NULL
        );
        """)

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS matches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            team1_name TEXT NOT NULL,
            team2_name TEXT NOT NULL,
            score_team1 INTEGER DEFAULT 0,
            score_team2 INTEGER DEFAULT 0,
            match_date TEXT
        );
        """)
        conn.commit()


def save_team(team: Team, db_path: str = DB_NAME) -> int:
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
        INSERT INTO teams (name, wins, draws, losses, goals_for, goals_against)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(name) DO UPDATE SET
            wins=excluded.wins,
            draws=excluded.draws,
            losses=excluded.losses,
            goals_for=excluded.goals_for,
            goals_against=excluded.goals_against;
        """, (team.name, team.wins, team.draws, team.losses, team.goals_for, team.goals_against))
        conn.commit()
        
        cursor.execute("SELECT id FROM teams WHERE name = ?", (team.name,))
        team_id = cursor.fetchone()[0]
        
        for player in team.players:
            save_player(player, team_id=team_id, db_path=db_path)
            
        return team_id


def save_player(player: Player, team_id: Optional[int] = None, db_path: str = DB_NAME):
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
        INSERT INTO players (name, position, matches, goals, assists, penalty_minutes, team_id)
        VALUES (?, ?, ?, ?, ?, ?, ?);
        """, (player.name, player.position, player.matches, player.goals, player.assists, player.penalty_minutes, team_id))
        conn.commit()


def save_match(match: Match, db_path: str = DB_NAME):
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
        INSERT INTO matches (team1_name, team2_name, score_team1, score_team2, match_date)
        VALUES (?, ?, ?, ?, ?);
        """, (match.team1_name, match.team2_name, match.score_team1, match.score_team2, match.date))
        conn.commit()

def export_to_docx(teams: List[Team], output_filename: str = "report.docx"):
    doc = docx.Document()
    doc.add_heading("Отчет по спортивной аналитике", level=0)

    for team in teams:
        doc.add_heading(f"Команда: {team.name}", level=1)
        doc.add_paragraph(team.get_summary())

        if team.players:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Имя"
            hdr_cells[1].text = "Позиция"
            hdr_cells[2].text = "Матчи"
            hdr_cells[3].text = "Голы"
            hdr_cells[4].text = "Передачи"

            for player in team.players:
                row_cells = table.add_row().cells
                row_cells[0].text = player.name
                row_cells[1].text = player.position
                row_cells[2].text = str(player.matches)
                row_cells[3].text = str(player.goals)
                row_cells[4].text = str(player.assists)
        
        doc.add_paragraph()

    doc.save(output_filename)


def export_to_xlsx(teams: List[Team], output_filename: str = "report.xlsx"):
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active  # type: ignore[assignment]
    if ws is None:
        ws = wb.create_sheet(title="Статистика игроков")
    else:
        ws.title = "Статистика игроков"

    headers = ["Команда", "Имя игрока", "Позиция", "Матчи", "Голы", "Передачи", "Штрафные минуты"]
    ws.append(headers)

    for team in teams:
        for player in team.players:
            ws.append([
                team.name,
                player.name,
                player.position,
                player.matches,
                player.goals,
                player.assists,
                player.penalty_minutes
            ])

    wb.save(output_filename)

def get_player_by_name(name: str, db_path: str = DB_NAME) -> Optional[Player]:
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT name, position, matches, goals, assists, penalty_minutes
            FROM players
            WHERE name = ?;
        """, (name,))
        row = cursor.fetchone()
        if row:
            return Player(
                name=row[0],
                position=row[1],
                matches=row[2],
                goals=row[3],
                assists=row[4],
                penalty_minutes=row[5]
            )
        return None


def get_all_teams_with_players(db_path: str = DB_NAME) -> List[Team]:
    teams_dict = {}
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        # Извлечение всех команд
        cursor.execute("SELECT id, name, wins, draws, losses, goals_for, goals_against FROM teams;")
        teams_rows = cursor.fetchall()
        for t_row in teams_rows:
            team_id, name, wins, draws, losses, g_for, g_against = t_row
            team = Team(
                name=name,
                wins=wins,
                draws=draws,
                losses=losses,
                goals_for=g_for,
                goals_against=g_against
            )
            teams_dict[team_id] = team

        # Извлечение всех игроков и привязка к командам
        cursor.execute("SELECT name, position, matches, goals, assists, penalty_minutes, team_id FROM players;")
        players_rows = cursor.fetchall()
        for p_row in players_rows:
            p_name, pos, matches, goals, assists, penalty, t_id = p_row
            player = Player(
                name=p_name,
                position=pos,
                matches=matches,
                goals=goals,
                assists=assists,
                penalty_minutes=penalty
            )
            if t_id in teams_dict:
                teams_dict[t_id].add_player(player)

    return list(teams_dict.values())

def get_or_create_team_id(team_name: str, db_path: str = DB_NAME):
    """Возвращает ID команды по названию. Если команда отсутствует в БД, создает новую."""
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM teams WHERE name = ?;", (team_name,))
        row = cursor.fetchone()
        if row:
            return row[0]
        
        # Если команда не найдена — создаем новую
        cursor.execute("INSERT INTO teams (name) VALUES (?);", (team_name,))
        conn.commit()
        return cursor.lastrowid